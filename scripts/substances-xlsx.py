# -*- coding: utf-8 -*-
"""
STEP1「化学式マッチ」の出題物質一覧を Excel で書き出す。

    npm run xlsx        （= python3 scripts/substances-xlsx.py）

src/data.js を直接読んで一覧を作るので、物質を足したり q を変えたりしたら
このスクリプトを流し直せば docs/ の Excel が最新になる。
openpyxl が要る（pip install openpyxl）。
"""
import json
import os
import re
import subprocess
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "docs", "STEP1_化学式マッチ_出題一覧.xlsx")

JP = "Yu Gothic"  # 日本語
EN = "Arial"      # 化学式・数字

HEAD_FILL = PatternFill("solid", fgColor="1F3864")
BAND_FILL = PatternFill("solid", fgColor="F2F5FA")
BASIC_FILL = PatternFill("solid", fgColor="E3F2FD")
CHAL_FILL = PatternFill("solid", fgColor="EDE7F6")
_thin = Side(style="thin", color="C5CBD6")
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


def load_substances():
    """src/data.js から物質データを取り出す（添字表記と単体判定は chem.js に任せる）"""
    script = """
import { SUBSTANCES } from './src/data.js';
import { formulaToUnicode, parseFormula } from './src/chem.js';
console.log(JSON.stringify(SUBSTANCES.map(function (s) {
  return {
    f: s.f,
    uni: formulaToUnicode(s.f),
    name: s.name,
    q: s.q,
    kind: Object.keys(parseFormula(s.f)).length === 1 ? '単体' : '化合物',
  };
})));
"""
    res = subprocess.run(
        ["node", "--input-type=module", "-e", script],
        cwd=ROOT, capture_output=True, text=True,
    )
    if res.returncode != 0:
        sys.exit("src/data.js を読めませんでした:\n" + res.stderr)
    return json.loads(res.stdout)


def write_sheet(ws, data, title, note):
    ws.sheet_view.showGridLines = False
    ws["A1"] = title
    ws["A1"].font = Font(name=JP, size=14, bold=True, color="1F3864")
    ws["A2"] = note
    ws["A2"].font = Font(name=JP, size=9, color="666666")
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[3].height = 6

    headers = ["No.", "レベル", "化学式", "化学式（添字）", "物質名", "単体／化合物"]
    hr = 4
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=hr, column=i, value=h)
        c.font = Font(name=JP, size=10, bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BORDER
    ws.row_dimensions[hr].height = 24

    for n, r in enumerate(data, start=1):
        row = hr + n
        level = {1: "基本", 2: "チャレンジ"}.get(r["q"], "出題なし")
        vals = [n, level, r["f"], r["uni"], r["name"], r["kind"]]
        for i, v in enumerate(vals, start=1):
            c = ws.cell(row=row, column=i, value=v)
            c.border = BORDER
            # 化学式の列だけ欧文フォント。日本語の列は和文フォント
            c.font = Font(name=EN if i in (1, 3, 4) else JP, size=11)
            c.alignment = Alignment(
                horizontal="center" if i in (1, 2, 6) else "left", vertical="center"
            )
            if n % 2 == 0:
                c.fill = BAND_FILL
        lv = ws.cell(row=row, column=2)
        if r["q"] == 1:
            lv.fill = BASIC_FILL
        elif r["q"] == 2:
            lv.fill = CHAL_FILL
        lv.font = Font(name=JP, size=10, bold=True)

    for i, w in enumerate([6, 12, 14, 16, 26, 14], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=hr + 1, column=1)
    return hr + len(data)


def main():
    rows = load_substances()
    main_rows = [r for r in rows if r["q"] in (1, 2)]
    extra_rows = [r for r in rows if r["q"] == 0]
    n_basic = sum(1 for r in main_rows if r["q"] == 1)
    n_chal = sum(1 for r in main_rows if r["q"] == 2)

    wb = Workbook()
    ws = wb.active
    ws.title = "STEP1 出題一覧"
    last = write_sheet(
        ws, main_rows,
        "STEP1「化学式マッチ」出題一覧",
        "物質名 ⇔ 化学式を4択で対応づけるモード。1回の挑戦で20問がランダムに出題されます。"
        "　／　基本 %d 物質・チャレンジ %d 物質（計 %d 物質）"
        % (n_basic, n_chal, len(main_rows)),
    )

    # 集計は数式で持つ（行を足し引きしても追従する）
    s = last + 2
    rng = "B5:B%d" % last
    for i, (label, formula) in enumerate([
        ("基本", '=COUNTIF(%s,"基本")' % rng),
        ("チャレンジ", '=COUNTIF(%s,"チャレンジ")' % rng),
        ("合計", "=SUM(C%d:C%d)" % (s, s + 1)),
    ]):
        ws.cell(row=s + i, column=2, value=label).font = Font(name=JP, size=10, bold=True)
        c = ws.cell(row=s + i, column=3, value=formula)
        c.font = Font(name=EN, size=10, bold=True)
        c.alignment = Alignment(horizontal="center")

    ws2 = wb.create_sheet("参考：出題されない物質")
    write_sheet(
        ws2, extra_rows,
        "参考：反応式には出てくるが STEP1 では出題されない物質",
        "STEP2〜FINAL の反応式やカードには登場しますが、化学式マッチの4択には出しません"
        "（src/data.js の q: 0）。",
    )

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    print("書き出しました: %s（基本 %d / チャレンジ %d / 参考 %d）"
          % (os.path.relpath(OUT, ROOT), n_basic, n_chal, len(extra_rows)))


if __name__ == "__main__":
    main()
